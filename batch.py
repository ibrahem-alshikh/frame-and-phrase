"""Batch-generate a whole set of images at once from an uploaded
filename -> text mapping (Excel or CSV), reusing the same saved template
settings for every image so the whole set stays visually consistent."""
import csv
import io
import os
from datetime import datetime

import openpyxl

from imaging import generate_image


def _stem(filename):
    root, _ext = os.path.splitext(os.path.basename(filename or ""))
    return root.strip()


def read_mapping(file_storage):
    """Read a two-column (filename, text) table from an uploaded .xlsx or
    .csv file. The first row is treated as a header and skipped. Returns
    {normalized_filename_stem: text}."""
    filename = (file_storage.filename or "").lower()

    if filename.endswith(".csv"):
        content = file_storage.read().decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(content)))
    else:
        workbook = openpyxl.load_workbook(file_storage, data_only=True)
        sheet = workbook.active
        rows = [[cell.value for cell in row] for row in sheet.iter_rows()]

    mapping = {}
    for row in rows[1:]:
        if not row or row[0] is None:
            continue
        key = _stem(str(row[0])).casefold()
        text = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        if key and text:
            mapping[key] = text
    return mapping


def downloads_dir():
    return os.path.join(os.path.expanduser("~"), "Downloads")


def _new_batch_dir_name():
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"نتائج_المجموعة_{stamp}"


def run_batch(image_files, mapping, settings):
    """image_files: list of werkzeug FileStorage. mapping: {stem: text}
    (as returned by read_mapping). Generates every matched pair with the
    same `settings` and writes each result to a fresh folder in Downloads,
    named after its source image. Returns a report dict.

    The output folder is only created once the first image actually
    succeeds, so a batch that fails entirely doesn't leave empty clutter
    behind in the user's Downloads folder.
    """
    out_dir = os.path.join(downloads_dir(), _new_batch_dir_name())
    dir_created = False
    results = []
    matched_keys = set()

    for f in image_files:
        stem = _stem(f.filename)
        text = mapping.get(stem.casefold())
        if text is None:
            results.append({"file": f.filename, "status": "missing_text"})
            continue
        matched_keys.add(stem.casefold())
        try:
            img = generate_image(f, text, settings)
            if not dir_created:
                os.makedirs(out_dir, exist_ok=True)
                dir_created = True
            out_path = os.path.join(out_dir, f"{stem}.png")
            img.save(out_path, format="PNG")
            results.append({"file": f.filename, "status": "ok", "text": text})
        except Exception as e:
            results.append({"file": f.filename, "status": "error", "message": str(e)})

    unmatched_texts = [key for key in mapping if key not in matched_keys]

    return {
        "output_dir": out_dir if dir_created else None,
        "total": len(image_files),
        "ok": sum(1 for r in results if r["status"] == "ok"),
        "results": results,
        "unmatched_texts": unmatched_texts,
    }


def build_template_workbook():
    """An empty starter .xlsx (header + one example row) users can fill in
    and re-upload, matching the exact column order read_mapping expects."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "الكلمات"
    sheet.append(["اسم الملف", "الكلمات"])
    sheet.append(["Q004S002", "جبال وأودية مكة"])
    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 40
    buf = io.BytesIO()
    workbook.save(buf)
    buf.seek(0)
    return buf
